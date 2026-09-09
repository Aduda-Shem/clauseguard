import io

import pytest
from docx import Document

from reviews.services.parsing import UnreadableFile, UnsupportedFileType, extract_text


def test_txt_extraction():
    assert extract_text("a.txt", b"hello world") == "hello world"


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedFileType):
        extract_text("a.xyz", b"whatever")


def test_docx_extraction_round_trip():
    doc = Document()
    doc.add_paragraph("Clause one.")
    doc.add_paragraph("Clause two.")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("contract.docx", buf.getvalue())
    assert "Clause one." in text
    assert "Clause two." in text


def test_corrupted_docx_raises_unreadable():
    with pytest.raises(UnreadableFile):
        extract_text("contract.docx", b"not a real docx file")


def test_corrupted_pdf_raises_unreadable():
    with pytest.raises(UnreadableFile):
        extract_text("contract.pdf", b"not a real pdf file")
